"""文档加载器 - 支持多种文件格式"""
import os
from pathlib import Path

# 文件扩展名 → MIME 类型白名单（批次D·11d：图片文档 PNG/JPG）
ALLOWED_MIME = {
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".pdf": "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".csv": "text/csv",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
}


def validate_file(file_path: str) -> bool:
    """校验文件扩展名在允许列表中"""
    ext = Path(file_path).suffix.lower()
    if ext not in ALLOWED_MIME:
        return False
    # 检查文件头魔数（docx/xlsx 都是 zip 容器，共用 PK 头）
    try:
        with open(file_path, "rb") as f:
            header = f.read(8)
        if ext == ".pdf" and not header.startswith(b"%PDF"):
            return False
        if ext in (".docx", ".xlsx") and not header.startswith(b"PK\x03\x04"):
            return False
        # 批次D·11d：图片魔数（PNG 89 50 4E 47；JPG FF D8 FF）
        if ext == ".png" and not header.startswith(b"\x89PNG"):
            return False
        if ext in (".jpg", ".jpeg") and not header.startswith(b"\xff\xd8\xff"):
            return False
    except Exception:
        return False
    return True


def load_document(file_path: str) -> str:
    """根据文件类型加载文档内容，返回纯文本"""
    ext = Path(file_path).suffix.lower()

    if ext == ".txt":
        return _load_txt(file_path)
    elif ext == ".md":
        return _load_txt(file_path)
    elif ext == ".pdf":
        return _load_pdf(file_path)
    elif ext == ".docx":
        return _load_docx(file_path)
    elif ext == ".csv":
        return _load_csv(file_path)
    else:
        raise ValueError(f"不支持的文件格式: {ext}")


def _load_txt(file_path: str) -> str:
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read()


def _load_pdf(file_path: str) -> str:
    """使用 pypdf 加载 PDF"""
    from pypdf import PdfReader
    reader = PdfReader(file_path)
    texts = []
    for page in reader.pages:
        text = page.extract_text()
        if text:
            texts.append(text)
    return "\n\n".join(texts)


def _load_docx(file_path: str) -> str:
    """使用 python-docx 加载 Word 文档"""
    from docx import Document as DocxDocument
    doc = DocxDocument(file_path)
    texts = [p.text for p in doc.paragraphs if p.text.strip()]
    # 也提取表格内容
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text for cell in row.cells)
            texts.append(row_text)
    return "\n\n".join(texts)


def _load_csv(file_path: str) -> str:
    """加载 CSV 文件，转为结构化文本"""
    import csv
    texts = []
    with open(file_path, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        headers = reader.fieldnames
        if headers:
            texts.append("列：" + ", ".join(headers))
        for i, row in enumerate(reader):
            texts.append(f"行{i+1}: " + " | ".join(f"{k}: {v}" for k, v in row.items()))
    return "\n".join(texts)


def load_table(file_path: str, fallback_name: str | None = None) -> list[dict]:
    """表格模板专用：结构化加载表格文件（表头 + 数据行），供 chunk_table 使用。

    返回 [{"name": 表名, "headers": [列头...], "rows": [[单元格...], ...]}]。
    CSV 按首个 sheet 处理（表名取 fallback_name 或文件名）；xlsx 每个 sheet 一项。
    fallback_name（验收修复 A3）：系统存储文件用 UUID 命名，磁盘文件名是 UUID——
    表名必须用原始文件名，否则分块表头显示一串 UUID 而不是"请假类型对照表"。
    """
    ext = Path(file_path).suffix.lower()
    if ext == ".csv":
        return _load_csv_table(file_path, fallback_name=fallback_name)
    if ext == ".xlsx":
        return _load_xlsx_table(file_path)
    raise ValueError(f"表格模板不支持的文件格式: {ext}")


def _load_csv_table(file_path: str, fallback_name: str | None = None) -> list[dict]:
    import csv

    with open(file_path, "r", encoding="utf-8-sig") as f:  # utf-8-sig 兼容 Excel 导出带 BOM
        rows = list(csv.reader(f))
    if not rows:
        return []
    name = fallback_name or Path(file_path).stem
    # 去掉扩展名（原始文件名带 .csv 后缀）
    if name.lower().endswith(".csv"):
        name = name[:-4]
    return [{"name": name, "headers": rows[0], "rows": rows[1:]}]


def _load_xlsx_table(file_path: str) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(file_path, read_only=True, data_only=True)
    try:
        sheets = []
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                # 跳过整行全空（表格末尾常有空白行）
                if any(v is not None and str(v).strip() for v in row):
                    rows.append(list(row))
            if rows:
                sheets.append({
                    "name": ws.title,
                    "headers": rows[0],
                    "rows": rows[1:],
                })
        return sheets
    finally:
        wb.close()
